import os
import pandas as pd
from datetime import datetime, timezone, timedelta
from docx import Document
from docx.shared import Inches
from utils import format_int, format_percentage, get_token_name
from config import CURRENT_REPORT_NUMBER, DOCX_SAVE_PATH, XGB_PARAMS, RATE_MAP
import re

def build_param_string(base_params: dict, extra: dict = {}) -> str:
    """Trả về chuỗi khởi tạo XGBClassifier(...) đẹp mắt để đưa vào báo cáo"""
    params = {**base_params, **(extra or {})}
    lines = [f"    {k}={v}," for k, v in params.items()]
    joined = "\n".join(lines)
    return f"xgb_model = XGBClassifier(\n{joined}\n)"

# --- tiện ích lấy Rate từ Y_TYPE ---
def get_rate_from_y_type(y_type: str, default="0.08") -> str:
    """Trả về chuỗi rate dựa trên hậu tố rXXX trong y_type"""
    m = re.search(r"r(\d+)", y_type)
    if not m:
        return default
    key = int(m.group(1))
    return str(RATE_MAP.get(key, default))

class ReportGenerator:
    def __init__(self, data_path, test_new_path):
        self.data_path = data_path
        self.test_new_path = test_new_path
        # Thay thế file counter bằng biến đơn giản
        self.current_report_number = CURRENT_REPORT_NUMBER
    
    
        
    def generate_word_report(self, results, output_dir):
        """Generate Word document report"""
        # Get auto-incrementing counter
        self.current_report_number += 1
        report_counter = self.current_report_number
        
        doc = Document()
        doc.add_heading("REPORT", level=1)
        doc.add_paragraph("")

        Y_TYPE = results['Y_TYPE']
        threshold = results['threshold']
        features = results['features']
        rate_value = get_rate_from_y_type(Y_TYPE)      # <-- dùng map

        # ------------------ MỤC 1 ------------------
        doc.add_heading("1. Thông tin tổng quan", level=2)

        # Get token name from file path
        token_name = get_token_name(self.data_path)

        info_1 = [
            ("Tên mô hình", "XGBoost"),
            ("Pair token", token_name),
            ("Dự đoán MinLow/MaxHigh", "MinLow" if "minlow" in Y_TYPE else "MaxHigh"),
            ("Threashold", str(threshold)),
            ("Rate", rate_value),                      # <-- thay chuỗi if-else dài
            ("Time frame prediction", "120M"),
            ("Có cân bằng dữ liệu không", "Có"),
            ("Exchange", "Binance"),
        ]
        table_1 = doc.add_table(rows=1, cols=2)
        table_1.style = 'Table Grid'
        table_1.rows[0].cells[0].text = "Thông tin"
        table_1.rows[0].cells[1].text = "Giá trị"
        for label_text, value in info_1:
            row = table_1.add_row().cells
            row[0].text = label_text
            row[1].text = value

        # ------------------ MỤC 1.2: Thống kê tập train ------------------
        doc.add_paragraph("\nTập dữ liệu train", style='Heading 3')
        
        # Get counts from cached stats to avoid reloading large CSV
        full_stats = results.get('full_dataset_stats', {})
        label_0_full = full_stats.get('count_0', 0)
        label_1_full = full_stats.get('count_1', 0)
        total_full = full_stats.get('total_count', 0)
        
        # Train set counts
        label_0_train = (results['y_train'] == 0).sum()
        label_1_train = (results['y_train'] == 1).sum()
        total_train = len(results['y_train'])
        pct_1_train = label_1_train / total_train * 100
        pct_0_train = label_0_train / total_train * 100
        
        train_info = [
            ("Tên bộ dữ liệu", os.path.basename(self.data_path)),
            ("Tỉ lệ dữ liệu train/ test", "70 / 30"),
            ("Các thuộc tính đầu vào", "\n".join([f"'{f}'" for f in features])),
            ("Số lượng nhãn 0", format_int(label_0_full)),
            ("Số lượng nhãn 1", format_int(label_1_full)),
            ("Số lượng bản ghi", format_int(total_full)),
            ("Số lượng nhãn 1 tập train", format_int(label_1_train)),
            ("Tỉ lệ % nhãn 1 tập train", f"{pct_1_train:.1f}%"),
            ("Số lượng nhãn 0 tập train", format_int(label_0_train)),
            ("Tỉ lệ % nhãn 0 tập train", f"{pct_0_train:.1f}%")
        ]
        table_train = doc.add_table(rows=1, cols=2)
        table_train.style = 'Table Grid'
        table_train.rows[0].cells[0].text = "Thông tin"
        table_train.rows[0].cells[1].text = "Giá trị"
        for label_text, value in train_info:
            row = table_train.add_row().cells
            row[0].text = label_text
            row[1].text = value

        # ------------------ MỤC 1.3: Dữ liệu mới ------------------
        doc.add_paragraph("\nTập dữ liệu Test (nằm ngoài tập dữ liệu train)", style='Heading 3')
        
        new_data_info = results["new_data"]
        test_info = [
            ("Tên bộ dữ liệu tập test", os.path.basename(self.test_new_path)),
            ("Số lượng bản ghi", format_int(new_data_info['count_rows_new'])),
            ("Số lượng nhãn 0 tập test", format_int(new_data_info['count_0'])),
            ("Tỉ lệ % nhãn 0 tập test", f"{new_data_info['pct_0']:.2f}%"),
            ("Số lượng nhãn 1 tập test", format_int(new_data_info['count_1'])),
            ("Tỉ lệ % nhãn 1 tập test", f"{new_data_info['pct_1']:.2f}%")
        ]
        table_test = doc.add_table(rows=1, cols=2)
        table_test.style = 'Table Grid'
        table_test.rows[0].cells[0].text = "Thông tin"
        table_test.rows[0].cells[1].text = "Giá trị"
        for label_text, value in test_info:
            row = table_test.add_row().cells
            row[0].text = label_text
            row[1].text = value

        # ------------------ MỤC 2 ------------------
        doc.add_heading("2. Bộ tham số đầu vào", level=2)
        extra = {"scale_pos_weight": f"{results['scale_pos_weight']:.2f}"}
        param_str = build_param_string(XGB_PARAMS, extra)
        doc.add_paragraph(param_str)

        # ------------------ MỤC 3 ------------------
        doc.add_heading("3. Kết quả chạy model", level=2)
        
        # Calculate accuracies using stored results
        train_accuracy = results['xgb_model'].score(results['X_train'], results['y_train']) * 100
        test_accuracy = results['accuracy_test'] * 100
        
        # Format as separate paragraphs like in the example
        doc.add_paragraph(f"Thời gian chạy\n{results['elapsed']:.1f}s")
        doc.add_paragraph(f"Train accuracy\n{train_accuracy:.2f}%")
        doc.add_paragraph(f"Test accuracy\n{test_accuracy:.2f}%")

        # Classification Report Test table
        test_report = results['report_test']
        
        table_report = doc.add_table(rows=1, cols=5)
        table_report.style = 'Table Grid'
        hdr_cells = table_report.rows[0].cells
        hdr_cells[0].text = "Label"
        hdr_cells[1].text = "Precision"
        hdr_cells[2].text = "Recall"
        hdr_cells[3].text = "F1-score"
        hdr_cells[4].text = "Support"
        
        # Add rows in specific order: 0, 1, accuracy, macro avg, weighted avg
        rows_data = []
        
        # Add class 0 and 1 first
        for class_label in ['0', '1']:
            if class_label in test_report:
                metrics = test_report[class_label]
                rows_data.append([
                    class_label,
                    format_percentage(metrics.get('precision', 0)),
                    format_percentage(metrics.get('recall', 0)),
                    format_percentage(metrics.get('f1-score', 0)),
                    format_int(metrics.get('support', 0))
                ])
        
        # Add accuracy row
        if 'accuracy' in test_report:
            rows_data.append([
                'Accuracy',
                '-',
                '-',
                format_percentage(test_report['accuracy']),
                format_int(len(results['y_test']))
            ])
        
        # Add macro avg
        if 'macro avg' in test_report:
            metrics = test_report['macro avg']
            rows_data.append([
                'Macro avg',
                format_percentage(metrics.get('precision', 0)),
                format_percentage(metrics.get('recall', 0)),
                format_percentage(metrics.get('f1-score', 0)),
                format_int(metrics.get('support', 0))
            ])
        
        # Add weighted avg
        if 'weighted avg' in test_report:
            metrics = test_report['weighted avg']
            rows_data.append([
                'Weighted avg',
                format_percentage(metrics.get('precision', 0)),
                format_percentage(metrics.get('recall', 0)),
                format_percentage(metrics.get('f1-score', 0)),
                format_int(metrics.get('support', 0))
            ])
        
        # Add all rows to table
        for row_data in rows_data:
            cells = table_report.add_row().cells
            for i, text in enumerate(row_data):
                cells[i].text = text
        
        # ------------------ MỤC 4 ------------------
        doc.add_heading("4. Các biểu đồ kết quả", level=2)
        for img in ["prediction_bar_test.png", "confusion_matrix_test.png", "feature_importance.png"]:
            path = os.path.join(output_dir, img)
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(5.5))
                doc.add_paragraph("")
        
        # ------------------ MỤC 5 & 6: New data results ------------------
        if 'report_new' in results:
            doc.add_heading("5. Kết quả chạy với bộ dữ liệu test ngoài", level=2)
            
            new_report = results['report_new']
            table_new = doc.add_table(rows=1, cols=5)
            table_new.style = 'Table Grid'
            hdr_cells_new = table_new.rows[0].cells
            hdr_cells_new[0].text = "Label"
            hdr_cells_new[1].text = "Precision"
            hdr_cells_new[2].text = "Recall"
            hdr_cells_new[3].text = "F1-score"
            hdr_cells_new[4].text = "Support"
            
            # Same order for new data table
            rows_new_data = []
            
            # Add class 0 and 1 first
            for class_label in ['0', '1']:
                if class_label in new_report:
                    metrics = new_report[class_label]
                    rows_new_data.append([
                        class_label,
                        format_percentage(metrics.get('precision', 0)),
                        format_percentage(metrics.get('recall', 0)),
                        format_percentage(metrics.get('f1-score', 0)),
                        format_int(metrics.get('support', 0))
                    ])
            
            # Add accuracy row
            if 'accuracy' in new_report:
                rows_new_data.append([
                    'Accuracy',
                    '-',
                    '-',
                    format_percentage(new_report['accuracy']),
                    format_int(new_data_info['count_rows_new'])
                ])
            
            # Add macro avg
            if 'macro avg' in new_report:
                metrics = new_report['macro avg']
                rows_new_data.append([
                    'Macro avg',
                    format_percentage(metrics.get('precision', 0)),
                    format_percentage(metrics.get('recall', 0)),
                    format_percentage(metrics.get('f1-score', 0)),
                    format_int(metrics.get('support', 0))
                ])
            
            # Add weighted avg
            if 'weighted avg' in new_report:
                metrics = new_report['weighted avg']
                rows_new_data.append([
                    'Weighted avg',
                    format_percentage(metrics.get('precision', 0)),
                    format_percentage(metrics.get('recall', 0)),
                    format_percentage(metrics.get('f1-score', 0)),
                    format_int(metrics.get('support', 0))
                ])
            
            # Add all rows to new table
            for row_data in rows_new_data:
                cells = table_new.add_row().cells
                for i, text in enumerate(row_data):
                    cells[i].text = text

            # ------------------ MỤC 6 ------------------
            doc.add_heading("6. Biểu đồ kết quả với tệp dữ liệu mới", level=2)
            for img in ["prediction_bar_new.png", "confusion_matrix_new.png"]:
                path = os.path.join(output_dir, img)
                if os.path.exists(path):
                    doc.add_picture(path, width=Inches(5.5))
                    doc.add_paragraph("")

        # Create advanced filename
        vietnam_tz = timezone(timedelta(hours=7))
        now_vn = datetime.now(vietnam_tz)
        now_str = now_vn.strftime("%d-%m-%Y")
        
        counter_str = str(report_counter).zfill(5)
        report_filename = f"Report-{counter_str}-Hiep-{token_name}USDT-XGB-{now_str}.docx"

        # Tạo thư mục custom nếu chưa có
        os.makedirs(DOCX_SAVE_PATH, exist_ok=True)

        # Lưu Word vào thư mục custom
        report_path = os.path.join(DOCX_SAVE_PATH, report_filename)
        doc.save(report_path)
        print(f"📄 Đã tạo báo cáo Word #{report_counter}: {report_path}")
        
        return report_path